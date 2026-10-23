"""Save the date 20 ans EDLC dans le style de la plaquette mécénat (A4), éléments natifs importables dans Canva."""
# Usage : python3 build_plaquette.py illustrations EDLC_Plaquette_mecenat.pptx
import sys
from PIL import Image
from pptx import Presentation
from pptx.util import Pt, Mm, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.oxml.ns import qn

IMG = sys.argv[1]
OUT = sys.argv[2]

# Charte EDLC
CORAL = RGBColor(0xEF, 0x3F, 0x34)
ORANGE = RGBColor(0xF8, 0x94, 0x10)
YELLOW = RGBColor(0xFF, 0xCB, 0x1F)
BLUE = RGBColor(0x26, 0x40, 0x8F)
GREEN = RGBColor(0x36, 0xB9, 0x95)
PINK = RGBColor(0xF5, 0x97, 0xB8)
INK = RGBColor(0x2B, 0x2B, 0x2B)
GREY = RGBColor(0x6B, 0x6B, 0x6B)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
# teintes claires dérivées des couleurs de la charte
CORAL_T = RGBColor(0xFD, 0xE8, 0xE6)
BLUE_T = RGBColor(0xE6, 0xEA, 0xF5)
YELLOW_T = RGBColor(0xFF, 0xF5, 0xD2)
ORANGE_T = RGBColor(0xFE, 0xEF, 0xDB)

TITLE_FONT = "Poppins"
BODY_FONT = "Open Sans"

W, H = 595.3, 841.9  # A4 en points
M = 36  # marge
CW = W - 2 * M

prs = Presentation()
prs.slide_width = Mm(210)
prs.slide_height = Mm(297)
blank = prs.slide_layouts[6]


def no_line(shp):
    shp.line.fill.background()


def rect(sl, x, y, w, h, fill, radius=None, line=None, lw=0.75, name=None):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    s = sl.shapes.add_shape(kind, Pt(x), Pt(y), Pt(w), Pt(h))
    if radius:
        s.adjustments[0] = min(0.5, radius / min(w, h))
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    if line is None:
        no_line(s)
    else:
        s.line.color.rgb = line
        s.line.width = Pt(lw)
    s.shadow.inherit = False
    s.text_frame.text = ""
    if name:
        s.name = name
    return s


def oval(sl, x, y, d, fill, name=None):
    s = sl.shapes.add_shape(MSO_SHAPE.OVAL, Pt(x), Pt(y), Pt(d), Pt(d))
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    no_line(s)
    s.shadow.inherit = False
    if name:
        s.name = name
    return s


def hline(sl, x, y, w, color, weight=0.75, name=None):
    c = sl.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Pt(x), Pt(y), Pt(x + w), Pt(y))
    c.line.color.rgb = color
    c.line.width = Pt(weight)
    if name:
        c.name = name
    return c


def text(sl, x, y, w, h, paras, size=9, color=INK, font=BODY_FONT, bold=False, italic=False,
         align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, spacing=1.2, after=0, name=None, spc=None):
    """paras: str | list de paragraphes ; un paragraphe = str | list de runs (texte, dict d'options)."""
    tb = sl.shapes.add_textbox(Pt(x), Pt(y), Pt(w), Pt(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    if isinstance(paras, str):
        paras = [paras]
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = spacing
        p.space_after = Pt(after)
        runs = [(para, {})] if isinstance(para, str) else [(r, {}) if isinstance(r, str) else r for r in para]
        for t, o in runs:
            r = p.add_run()
            r.text = t
            f = r.font
            f.name = o.get("font", font)
            f.size = Pt(o.get("size", size))
            f.bold = o.get("bold", bold)
            f.italic = o.get("italic", italic)
            f.color.rgb = o.get("color", color)
            s_ = o.get("spc", spc)
            if s_ is not None:
                r._r.get_or_add_rPr().set("spc", str(int(s_ * 100)))
            if "link" in o:
                r.hyperlink.address = o["link"]
                f.color.rgb = o.get("color", color)
                f.underline = True
    if name:
        tb.name = name
    return tb


def picture(sl, path, x, y, w=None, h=None, name=None):
    im = Image.open(path)
    bbox = im.getchannel("A").getbbox()
    if bbox and bbox != (0, 0) + im.size:
        im = im.crop(bbox)
        path = path.replace(".png", "_crop.png")
        im.save(path)
    iw, ih = im.size
    if w and not h:
        h = w * ih / iw
    if h and not w:
        w = h * iw / ih
    p = sl.shapes.add_picture(path, Pt(x), Pt(y), Pt(w), Pt(h))
    if name:
        p.name = name
    return p, w, h


def group(sl, shapes, name):
    g = sl.shapes.add_group_shape(shapes)
    g.name = name
    return g


def section_title(sl, x, y, label, size=15, color=CORAL, bar=CORAL, w=CW, name=None):
    b = rect(sl, x, y + 3, 4.5, size + 2, bar, radius=2.2, name=(name or label) + " – repère")
    t = text(sl, x + 11, y, w - 11, size * 1.5, label, size=size, color=color, font=TITLE_FONT,
             bold=True, spacing=1.0, name=(name or label))
    return b, t


def sub_title(sl, x, y, label, w=CW, color=BLUE, size=11.5):
    return text(sl, x, y, w, size * 1.5, label, size=size, color=color, font=TITLE_FONT, bold=True,
                spacing=1.0, name=label)


def chrome_top(sl):
    """Bandeau haut et en-tête (texte de l’en-tête du Word)."""
    rect(sl, 0, 0, W, 7, CORAL, name="Bandeau haut corail")
    text(sl, W - M - 300, 15, 300, 10, "ENTREPRENEURS DE LA CITÉ  |  FONDATION RECONNUE D’UTILITÉ PUBLIQUE",
         size=6.5, color=GREY, align=PP_ALIGN.RIGHT, spc=0.4, name="En-tête")


def chrome_bottom(sl):
    """Pied de page (texte du pied de page du Word)."""
    hline(sl, M, H - 34, CW, CORAL, 0.75, name="Pied de page – filet")
    text(sl, M, H - 27, CW, 10, [[("ENTREPRENEURS DE LA CITÉ", {"bold": True, "color": CORAL}),
                                  ("  •  Prévenir, protéger, pérenniser", {})]],
         size=6.8, color=GREY, align=PP_ALIGN.CENTER, spc=0.3, name="Pied de page")


def styled_table(sl, x, y, col_widths, row_heights, rows, style, name):
    """Tableau natif. rows : liste de lignes ; cellule = list de paragraphes (list de runs (texte, opts)).
    style(r, j) -> (fond, couleur texte, taille, gras, alignement)."""
    gt = sl.shapes.add_table(len(rows), len(col_widths), Pt(x), Pt(y), Pt(sum(col_widths)), Pt(sum(row_heights)))
    gt.name = name
    tbl = gt.table
    tbl._tbl.tblPr.set("firstRow", "0"); tbl._tbl.tblPr.set("bandRow", "0")
    for j, wcol in enumerate(col_widths):
        tbl.columns[j].width = Pt(wcol)
    for r, hrow in enumerate(row_heights):
        tbl.rows[r].height = Pt(hrow)
    for r, row in enumerate(rows):
        for j, paras in enumerate(row):
            fill, color, size, bold, align = style(r, j)
            cell = tbl.cell(r, j)
            cell.margin_left = cell.margin_right = Pt(6)
            cell.margin_top = cell.margin_bottom = Pt(2)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid(); cell.fill.fore_color.rgb = fill
            tf = cell.text_frame
            tf.word_wrap = True
            for k, para in enumerate(paras):
                p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
                p.alignment = align
                p.line_spacing = 1.0
                runs = [(para, {})] if isinstance(para, str) else para
                for t, o in runs:
                    run = p.add_run(); run.text = t
                    f = run.font; f.name = BODY_FONT
                    f.size = Pt(o.get("size", size)); f.bold = o.get("bold", bold)
                    f.color.rgb = o.get("color", color)
            tcPr = cell._tc.get_or_add_tcPr()
            for edge in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
                ln = tcPr.makeelement(qn(edge), {"w": str(Pt(0.5)), "cap": "flat", "cmpd": "sng", "algn": "ctr"})
                sf = ln.makeelement(qn("a:solidFill"), {})
                sf.append(sf.makeelement(qn("a:srgbClr"), {"val": "D9DCE6"}))
                ln.append(sf)
                tcPr.append(ln)
    return gt


# ---------- pictogrammes vectoriels (formes natives) ----------
def preset(sl, kind, x, y, w, h, fill=None, line=None, lw=0.9, name=None, rot=0):
    sh = sl.shapes.add_shape(kind, Pt(x), Pt(y), Pt(w), Pt(h))
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid(); sh.fill.fore_color.rgb = fill
    if line is None:
        no_line(sh)
    else:
        sh.line.color.rgb = line; sh.line.width = Pt(lw)
    sh.rotation = rot
    if name:
        sh.name = name
    return sh


def icon_protect(sl, x, y, d, bg):
    """Bouclier + coche (formes natives, non groupées pour rester vectorielles dans Canva)."""
    oval(sl, x, y, d, bg, name="Picto protéger – fond")
    s = d / 30
    preset(sl, MSO_SHAPE.FLOWCHART_OFFPAGE_CONNECTOR, x + 8 * s, y + 6.5 * s, 14 * s, 17 * s, WHITE,
           name="Picto protéger – bouclier")
    # coche = deux barres arrondies pivotées
    preset(sl, MSO_SHAPE.ROUNDED_RECTANGLE, x + 10.2 * s, y + 13.4 * s, 5 * s, 2.2 * s, bg, rot=45,
           name="Picto protéger – coche (court)")
    preset(sl, MSO_SHAPE.ROUNDED_RECTANGLE, x + 12.6 * s, y + 11.8 * s, 8.5 * s, 2.2 * s, bg, rot=-50,
           name="Picto protéger – coche (long)")


def icon_prevent(sl, x, y, d, bg, fg):
    """Ampoule = prévention / formation."""
    oval(sl, x, y, d, bg, name="Picto prévenir – fond")
    s = d / 30
    oval(sl, x + 9 * s, y + 5 * s, 12 * s, fg, name="Picto prévenir – ampoule")
    rect(sl, x + 12 * s, y + 15.5 * s, 6 * s, 5 * s, fg, name="Picto prévenir – culot")
    rect(sl, x + 12 * s, y + 21.3 * s, 6 * s, 1.4 * s, fg, name="Picto prévenir – trait 1")
    rect(sl, x + 13 * s, y + 23.5 * s, 4 * s, 1.4 * s, fg, name="Picto prévenir – trait 2")


def icon_pin(sl, x, y, d, color, hole):
    preset(sl, MSO_SHAPE.FLOWCHART_MERGE, x + d * 0.14, y + d * 0.55, d * 0.72, d * 0.75, color,
           name="Picto adresse – pointe")
    oval(sl, x, y, d, color, name="Picto adresse – tête")
    oval(sl, x + d * 0.3, y + d * 0.3, d * 0.4, hole, name="Picto adresse – centre")


def icon_web(sl, x, y, d, color):
    preset(sl, MSO_SHAPE.OVAL, x, y, d, d, None, line=color, name="Picto web – globe")
    preset(sl, MSO_SHAPE.OVAL, x + d * 0.3, y, d * 0.4, d, None, line=color, name="Picto web – méridien")
    hline(sl, x, y + d / 2, d, color, 0.9, name="Picto web – équateur")


# =====================================================================
# Save the date – style plaquette mécénat (texte du save the date validé, inchangé)
# Le grand « 20 » et le QR code sont ajoutés ensuite directement dans Canva.
# =====================================================================
M = 34
CW = W - 2 * M
s = prs.slides.add_slide(blank)

picture(s, f"{IMG}/logo_edlc.png", M, 26, w=118, name="Logo EDLC")

# panneau du grand « 20 »
rect(s, M, 76, CW, 250, BLUE_T, radius=16, name="Bloc 20 – fond bleu pâle")

# titre
text(s, M, 344, CW, 36, "SAVE THE DATE", size=30, font=TITLE_FONT, bold=True, color=BLUE,
     align=PP_ALIGN.CENTER, spacing=1.0, name="Titre – save the date")
text(s, M, 382, CW, 30, "20 ANS D’ENGAGEMENT !", size=22, font=TITLE_FONT, bold=True, color=CORAL,
     align=PP_ALIGN.CENTER, spacing=1.0, name="Titre – 20 ans d’engagement")

# texte d’invitation
text(s, M + 28, 424, CW - 56, 80,
     "La Fondation Entrepreneurs de la Cité a le plaisir de vous inviter à sa soirée anniversaire : Retour sur "
     "20 ans d’engagement auprès des entrepreneurs vulnérables à travers l’assurance sociale et solidaire. Au "
     "programme : Rétrospective en images, invités de prestige et projection vers l’avenir... une soirée à ne pas "
     "manquer !", size=10.5, color=INK, align=PP_ALIGN.CENTER, spacing=1.35, name="Invitation – texte")

# cases date / heure (style chiffres clés)
ty, th, gap = 518, 92, 9
tw = (CW - gap) / 2
rect(s, M, ty, tw, th, BLUE_T, radius=12, name="Case date – fond")
text(s, M + 58, ty + 12, 120, 16, "Mardi", size=11, bold=True, color=BLUE, spc=1.2, name="Date – jour")
text(s, M + 54, ty + 26, 80, 56, "15", size=44, font=TITLE_FONT, bold=True, color=BLUE, spacing=1.0,
     name="Date – 15")
text(s, M + 128, ty + 46, 120, 26, "juin 2027", size=17, font=TITLE_FONT, bold=True, color=BLUE,
     spacing=1.0, name="Date – mois")
rect(s, M + tw + gap, ty, tw, th, CORAL_T, radius=12, name="Case heure – fond")
text(s, M + tw + gap, ty, tw, th, "18H30", size=34, font=TITLE_FONT, bold=True, color=CORAL,
     align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, spacing=1.0, name="Heure")

# case lieu
ly = ty + th + 9
rect(s, M, ly, CW, 44, ORANGE_T, radius=12, name="Case lieu – fond")
icon_pin(s, M + 92, ly + 14, 11, ORANGE, ORANGE_T)
text(s, M + 112, ly, CW - 130, 44, "Le Philanthro-Lab, 13-15 rue de la Bûcherie, 75005 Paris", size=11.5,
     color=BLUE, bold=True, anchor=MSO_ANCHOR.MIDDLE, name="Lieu")
text(s, M, ly + 54, CW, 16, "Soirée animée par Jean-Luc Gambey associé à Vovoxx Média", size=10,
     italic=True, color=GREY, align=PP_ALIGN.CENTER, name="Animation")

# bloc inscription (style « Engageons-nous ensemble »)
by = 700
rect(s, M, by, CW, 110, BLUE, radius=14, name="Inscription – fond bleu")
text(s, M + 24, by + 26, 330, 30, "Inscription", size=22, font=TITLE_FONT, bold=True, color=WHITE,
     spacing=1.0, name="Inscription – titre")
rect(s, M + 24, by + 58, 28, 2.5, CORAL, name="Inscription – trait")
tb = text(s, M + 24, by + 70, 330, 16, [[("Cliquez ", {}), ("ici", {}), (" ou scannez ce QR code :", {})]],
          size=11.5, color=WHITE, name="Inscription – texte")
tb.text_frame.paragraphs[0].runs[1].font.underline = True
rect(s, M + CW - 112, by + 11, 88, 88, WHITE, radius=10, name="QR – cadre blanc")

for el in list(s.shapes._spTree.iter(qn("p:style"))):
    el.getparent().remove(el)
prs.save(OUT)
print("ok", OUT)
