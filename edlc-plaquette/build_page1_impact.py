"""Page 1 « impact » de la plaquette EDLC (A4 portrait), éléments natifs importables dans Canva."""
# Usage : python3 build_plaquette.py illustrations EDLC_Plaquette_mecenat.pptx
import sys
from PIL import Image
from pptx import Presentation
from pptx.util import Pt, Mm, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.oxml.ns import qn

IMG = sys.argv[1]
OUT = sys.argv[2]

# Charte EDLC
CORAL = RGBColor(0xEF, 0x3F, 0x34)
ORANGE = RGBColor(0xF8, 0x94, 0x10)
YELLOW = RGBColor(0xFF, 0xCB, 0x1F)
BLUE = RGBColor(0x26, 0x40, 0x8F)
GREEN = RGBColor(0x36, 0xB9, 0x95)
PINK = RGBColor(0xF5, 0x97, 0xB8)
INK = RGBColor(0x2B, 0x2B, 0x2B)
GREY = RGBColor(0x6B, 0x6B, 0x6B)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
# teintes claires dérivées des couleurs de la charte
CORAL_T = RGBColor(0xFD, 0xE8, 0xE6)
BLUE_T = RGBColor(0xE6, 0xEA, 0xF5)
YELLOW_T = RGBColor(0xFF, 0xF5, 0xD2)
ORANGE_T = RGBColor(0xFE, 0xEF, 0xDB)

TITLE_FONT = "Poppins"
BODY_FONT = "Open Sans"

W, H = 595.3, 841.9  # A4 en points
M = 36  # marge
CW = W - 2 * M

prs = Presentation()
prs.slide_width = Mm(210)
prs.slide_height = Mm(297)
blank = prs.slide_layouts[6]


def no_line(shp):
    shp.line.fill.background()


def rect(sl, x, y, w, h, fill, radius=None, line=None, lw=0.75, name=None):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    s = sl.shapes.add_shape(kind, Pt(x), Pt(y), Pt(w), Pt(h))
    if radius:
        s.adjustments[0] = min(0.5, radius / min(w, h))
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    if line is None:
        no_line(s)
    else:
        s.line.color.rgb = line
        s.line.width = Pt(lw)
    s.shadow.inherit = False
    s.text_frame.text = ""
    if name:
        s.name = name
    return s


def oval(sl, x, y, d, fill, name=None):
    s = sl.shapes.add_shape(MSO_SHAPE.OVAL, Pt(x), Pt(y), Pt(d), Pt(d))
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    no_line(s)
    s.shadow.inherit = False
    if name:
        s.name = name
    return s


def hline(sl, x, y, w, color, weight=0.75, name=None):
    c = sl.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Pt(x), Pt(y), Pt(x + w), Pt(y))
    c.line.color.rgb = color
    c.line.width = Pt(weight)
    if name:
        c.name = name
    return c


def text(sl, x, y, w, h, paras, size=9, color=INK, font=BODY_FONT, bold=False, italic=False,
         align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, spacing=1.2, after=0, name=None, spc=None):
    """paras: str | list de paragraphes ; un paragraphe = str | list de runs (texte, dict d'options)."""
    tb = sl.shapes.add_textbox(Pt(x), Pt(y), Pt(w), Pt(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    if isinstance(paras, str):
        paras = [paras]
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = spacing
        p.space_after = Pt(after)
        runs = [(para, {})] if isinstance(para, str) else [(r, {}) if isinstance(r, str) else r for r in para]
        for t, o in runs:
            r = p.add_run()
            r.text = t
            f = r.font
            f.name = o.get("font", font)
            f.size = Pt(o.get("size", size))
            f.bold = o.get("bold", bold)
            f.italic = o.get("italic", italic)
            f.color.rgb = o.get("color", color)
            s_ = o.get("spc", spc)
            if s_ is not None:
                r._r.get_or_add_rPr().set("spc", str(int(s_ * 100)))
            if "link" in o:
                r.hyperlink.address = o["link"]
                f.color.rgb = o.get("color", color)
                f.underline = True
    if name:
        tb.name = name
    return tb


def picture(sl, path, x, y, w=None, h=None, name=None):
    im = Image.open(path)
    bbox = im.getchannel("A").getbbox()
    if bbox and bbox != (0, 0) + im.size:
        im = im.crop(bbox)
        path = path.replace(".png", "_crop.png")
        im.save(path)
    iw, ih = im.size
    if w and not h:
        h = w * ih / iw
    if h and not w:
        w = h * iw / ih
    p = sl.shapes.add_picture(path, Pt(x), Pt(y), Pt(w), Pt(h))
    if name:
        p.name = name
    return p, w, h


def group(sl, shapes, name):
    g = sl.shapes.add_group_shape(shapes)
    g.name = name
    return g


def section_title(sl, x, y, label, size=15, color=CORAL, bar=CORAL, w=CW, name=None):
    b = rect(sl, x, y + 3, 4.5, size + 2, bar, radius=2.2, name=(name or label) + " – repère")
    t = text(sl, x + 11, y, w - 11, size * 1.5, label, size=size, color=color, font=TITLE_FONT,
             bold=True, spacing=1.0, name=(name or label))
    return b, t


def sub_title(sl, x, y, label, w=CW, color=BLUE, size=11.5):
    return text(sl, x, y, w, size * 1.5, label, size=size, color=color, font=TITLE_FONT, bold=True,
                spacing=1.0, name=label)


def chrome_top(sl):
    """Bandeau haut et en-tête (texte de l’en-tête du Word)."""
    rect(sl, 0, 0, W, 7, CORAL, name="Bandeau haut corail")
    text(sl, W - M - 300, 15, 300, 10, "ENTREPRENEURS DE LA CITÉ  |  FONDATION RECONNUE D’UTILITÉ PUBLIQUE",
         size=6.5, color=GREY, align=PP_ALIGN.RIGHT, spc=0.4, name="En-tête")


def chrome_bottom(sl):
    """Pied de page (texte du pied de page du Word)."""
    hline(sl, M, H - 34, CW, CORAL, 0.75, name="Pied de page – filet")
    text(sl, M, H - 27, CW, 10, [[("ENTREPRENEURS DE LA CITÉ", {"bold": True, "color": CORAL}),
                                  ("  •  Prévenir, protéger, pérenniser", {})]],
         size=6.8, color=GREY, align=PP_ALIGN.CENTER, spc=0.3, name="Pied de page")


def styled_table(sl, x, y, col_widths, row_heights, rows, style, name):
    """Tableau natif. rows : liste de lignes ; cellule = list de paragraphes (list de runs (texte, opts)).
    style(r, j) -> (fond, couleur texte, taille, gras, alignement)."""
    gt = sl.shapes.add_table(len(rows), len(col_widths), Pt(x), Pt(y), Pt(sum(col_widths)), Pt(sum(row_heights)))
    gt.name = name
    tbl = gt.table
    tbl._tbl.tblPr.set("firstRow", "0"); tbl._tbl.tblPr.set("bandRow", "0")
    for j, wcol in enumerate(col_widths):
        tbl.columns[j].width = Pt(wcol)
    for r, hrow in enumerate(row_heights):
        tbl.rows[r].height = Pt(hrow)
    for r, row in enumerate(rows):
        for j, paras in enumerate(row):
            fill, color, size, bold, align = style(r, j)
            cell = tbl.cell(r, j)
            cell.margin_left = cell.margin_right = Pt(6)
            cell.margin_top = cell.margin_bottom = Pt(2)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid(); cell.fill.fore_color.rgb = fill
            tf = cell.text_frame
            tf.word_wrap = True
            for k, para in enumerate(paras):
                p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
                p.alignment = align
                p.line_spacing = 1.0
                runs = [(para, {})] if isinstance(para, str) else para
                for t, o in runs:
                    run = p.add_run(); run.text = t
                    f = run.font; f.name = BODY_FONT
                    f.size = Pt(o.get("size", size)); f.bold = o.get("bold", bold)
                    f.color.rgb = o.get("color", color)
            tcPr = cell._tc.get_or_add_tcPr()
            for edge in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
                ln = tcPr.makeelement(qn(edge), {"w": str(Pt(0.5)), "cap": "flat", "cmpd": "sng", "algn": "ctr"})
                sf = ln.makeelement(qn("a:solidFill"), {})
                sf.append(sf.makeelement(qn("a:srgbClr"), {"val": "D9DCE6"}))
                ln.append(sf)
                tcPr.append(ln)
    return gt


# ---------- pictogrammes vectoriels (formes natives) ----------
def preset(sl, kind, x, y, w, h, fill=None, line=None, lw=0.9, name=None, rot=0):
    sh = sl.shapes.add_shape(kind, Pt(x), Pt(y), Pt(w), Pt(h))
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid(); sh.fill.fore_color.rgb = fill
    if line is None:
        no_line(sh)
    else:
        sh.line.color.rgb = line; sh.line.width = Pt(lw)
    sh.rotation = rot
    if name:
        sh.name = name
    return sh


def icon_protect(sl, x, y, d, bg):
    """Bouclier + coche (formes natives, non groupées pour rester vectorielles dans Canva)."""
    oval(sl, x, y, d, bg, name="Picto protéger – fond")
    s = d / 30
    preset(sl, MSO_SHAPE.FLOWCHART_OFFPAGE_CONNECTOR, x + 8 * s, y + 6.5 * s, 14 * s, 17 * s, WHITE,
           name="Picto protéger – bouclier")
    # coche = deux barres arrondies pivotées
    preset(sl, MSO_SHAPE.ROUNDED_RECTANGLE, x + 10.2 * s, y + 13.4 * s, 5 * s, 2.2 * s, bg, rot=45,
           name="Picto protéger – coche (court)")
    preset(sl, MSO_SHAPE.ROUNDED_RECTANGLE, x + 12.6 * s, y + 11.8 * s, 8.5 * s, 2.2 * s, bg, rot=-50,
           name="Picto protéger – coche (long)")


def icon_prevent(sl, x, y, d, bg, fg):
    """Ampoule = prévention / formation."""
    oval(sl, x, y, d, bg, name="Picto prévenir – fond")
    s = d / 30
    oval(sl, x + 9 * s, y + 5 * s, 12 * s, fg, name="Picto prévenir – ampoule")
    rect(sl, x + 12 * s, y + 15.5 * s, 6 * s, 5 * s, fg, name="Picto prévenir – culot")
    rect(sl, x + 12 * s, y + 21.3 * s, 6 * s, 1.4 * s, fg, name="Picto prévenir – trait 1")
    rect(sl, x + 13 * s, y + 23.5 * s, 4 * s, 1.4 * s, fg, name="Picto prévenir – trait 2")


def icon_pin(sl, x, y, d, color, hole):
    preset(sl, MSO_SHAPE.FLOWCHART_MERGE, x + d * 0.14, y + d * 0.55, d * 0.72, d * 0.75, color,
           name="Picto adresse – pointe")
    oval(sl, x, y, d, color, name="Picto adresse – tête")
    oval(sl, x + d * 0.3, y + d * 0.3, d * 0.4, hole, name="Picto adresse – centre")


def icon_web(sl, x, y, d, color):
    preset(sl, MSO_SHAPE.OVAL, x, y, d, d, None, line=color, name="Picto web – globe")
    preset(sl, MSO_SHAPE.OVAL, x + d * 0.3, y, d * 0.4, d, None, line=color, name="Picto web – méridien")
    hline(sl, x, y + d / 2, d, color, 0.9, name="Picto web – équateur")


# =====================================================================
# PAGE 1 – version « impact » (texte fourni par EDLC, inchangé)
# =====================================================================
M = 34
CW = W - 2 * M
NAVY_T = RGBColor(0x33, 0x50, 0xA3)
GREEN_T = RGBColor(0xE1, 0xF5, 0xEF)
B = {"bold": True}

s = prs.slides.add_slide(blank)


def icon_info(sl, x, y, d, bg, fg):
    oval(sl, x, y, d, bg, name="Picto prévenir – fond")
    k = d / 30
    oval(sl, x + 13.2 * k, y + 7 * k, 3.6 * k, fg, name="Picto prévenir – point")
    rect(sl, x + 13.4 * k, y + 12.5 * k, 3.2 * k, 10.5 * k, fg, name="Picto prévenir – barre")


def icon_megaphone(sl, x, y, d, bg, fg):
    oval(sl, x, y, d, bg, name="Picto mobiliser – fond")
    k = d / 30
    preset(sl, MSO_SHAPE.TRAPEZOID, x + 9 * k, y + 9 * k, 12 * k, 10 * k, fg, rot=-90,
           name="Picto mobiliser – pavillon")
    rect(sl, x + 7.5 * k, y + 11.5 * k, 4 * k, 5 * k, fg, name="Picto mobiliser – arrière")
    rect(sl, x + 9 * k, y + 16 * k, 2.4 * k, 6 * k, fg, name="Picto mobiliser – poignée")
    for i, (dy, rot) in enumerate(((7.5, -30), (13.3, 0), (19.1, 30))):
        preset(sl, MSO_SHAPE.ROUNDED_RECTANGLE, x + 21.5 * k, y + dy * k, 4 * k, 1.4 * k, fg, rot=rot,
               name=f"Picto mobiliser – onde {i + 1}")


# ---------- en-tête ----------
rect(s, 0, 0, W, 6, CORAL, name="Bandeau haut corail")
picture(s, f"{IMG}/logo_edlc.png", M, 20, w=118, name="Logo EDLC")

# ---------- héro ----------
hero_y, hero_h = 62, 106
rect(s, M, hero_y, CW, hero_h, BLUE_T, radius=14, name="Héro – fond")
rect(s, M + 16, hero_y + 20, 4, hero_h - 40, CORAL, name="Héro – filet corail")
oval(s, M + CW - 118, hero_y + 9, 88, WHITE, name="Héro – cercle blanc")
picture(s, f"{IMG}/symbole_edlc.png", M + CW - 74 - 26, hero_y + 9 + 44 - 30, h=60,
        name="Pictogramme EDLC")
text(s, M + 30, hero_y + 14, 370, 62, ["L’assurance sociale et solidaire", "au service de l’entrepreneuriat"],
     size=21, font=TITLE_FONT, bold=True, color=BLUE, spacing=1.0, name="Titre principal")
text(s, M + 30, hero_y + 68, 370, 34,
     ["Devenez mécène d’une économie où chacun peut", "entreprendre et faire durer son activité !"],
     size=11.5, font=TITLE_FONT, bold=True, color=CORAL, spacing=1.05, name="Sous-titre")

# ---------- introduction ----------
y = hero_y + hero_h + 14
rect(s, M, y + 2, 3, 74, BLUE, name="Introduction – filet")
text(s, M + 12, y, CW - 12, 80, [
    [("Les créateurs d’entreprise en situation de précarité ou de handicap peuvent rencontrer des ", {}),
     ("difficultés à accéder aux assurances professionnelles classiques", B),
     (". Sans protection adaptée, un accident, un litige ou un arrêt de travail peut compromettre l’activité "
      "qu’ils ont créée.", {})],
    [("Depuis 2007, la Fondation Entrepreneurs de la Cité (EDLC) ouvre ", {}),
     ("l’accès à une assurance adaptée aux créateurs d’entreprise en situation de vulnérabilité", B),
     (". Nous agissons là où les offres classiques ne répondent pas aux besoins : au croisement de "
      "l’inclusion, de l’emploi et de la pérennité économique.", {})]],
    size=8.6, spacing=1.25, after=5, name="Introduction")

# ---------- chiffres clés : bandeau bleu ----------
y += 88
fig_h = 104
rect(s, M, y, CW, fig_h, BLUE, radius=14, name="Chiffres clés – fond")
cw3 = CW / 3
figs = [("3 362", ["entrepreneurs assurés", "en 2025"]),
        ("2 544", ["personnes formées aux risques", "en 2025"]),
        ("89 %", ["de survie après sinistre"]),
        ("11 741", ["entrepreneurs assurés depuis 2007"]),
        ("38 816", ["personnes formées aux risques", "depuis 2007"]),
        ("350+", ["activités couvertes"])]
num_colors = [YELLOW, WHITE, YELLOW]
for i, (num, lab) in enumerate(figs):
    r, c = divmod(i, 3)
    x = M + c * cw3
    yy = y + 8 + r * 48
    text(s, x, yy, cw3, 22, num, size=18, font=TITLE_FONT, bold=True, color=YELLOW if c != 1 else WHITE,
         align=PP_ALIGN.CENTER, spacing=1.0, name=f"Chiffre {num}")
    text(s, x + 6, yy + 22, cw3 - 12, 22, lab, size=7.4, color=WHITE, align=PP_ALIGN.CENTER, spacing=1.05,
         name=f"Chiffre {num} – légende")
for c in (1, 2):
    preset(s, MSO_SHAPE.RECTANGLE, M + c * cw3, y + 14, 0.6, fig_h - 28, NAVY_T,
           name=f"Chiffres clés – séparateur {c}")
rect(s, M + 22, y + fig_h / 2, CW - 44, 0.6, NAVY_T, name="Chiffres clés – séparateur horizontal")

# ---------- modèle ----------
y += fig_h + 14
text(s, M, y, CW, 20, "Un modèle qui associe prévention, protection et mobilisation", size=14,
     font=TITLE_FONT, bold=True, color=BLUE, spacing=1.0, name="Titre – modèle")
y += 24
gap = 9
cw = (CW - 2 * gap) / 3
card_h = 131
cards = [
    ("PRÉVENIR", BLUE, BLUE_T, "info",
     "Des webinaires hebdomadaires, des formations et des interventions auprès des réseaux "
     "d’accompagnement pour aider les entrepreneurs à identifier les risques et à choisir les protections "
     "adaptées."),
    ("PROTÉGER", CORAL, CORAL_T, "shield",
     "Des solutions assurantielles à tarifs solidaires, couvrant notamment la multirisque professionnelle, "
     "la responsabilité civile professionnelle, la santé, la prévoyance et la responsabilité civile "
     "décennale. L’assistance juridique est offerte."),
    ("MOBILISER", ORANGE, ORANGE_T, "mega",
     "Démonstration par la preuve de l'utilité sociale de l'assurance inclusive et son impact sur la "
     "société et la vitalité économique des territoires."),
]
for i, (lab, c, t_, ic, body) in enumerate(cards):
    x = M + i * (cw + gap)
    rect(s, x, y, cw, card_h, t_, radius=12, name=f"Carte {lab} – fond")
    rect(s, x + 14, y, cw - 28, 3.5, c, name=f"Carte {lab} – accent")
    if ic == "info":
        icon_info(s, x + 12, y + 14, 26, c, WHITE)
    elif ic == "shield":
        icon_protect(s, x + 12, y + 14, 26, c)
    else:
        icon_megaphone(s, x + 12, y + 14, 26, c, WHITE)
    text(s, x + 45, y + 14, cw - 52, 26, lab, size=11, font=TITLE_FONT, bold=True, color=c, spc=0.8,
         anchor=MSO_ANCHOR.MIDDLE, spacing=1.0, name=f"Carte {lab} – titre")
    text(s, x + 12, y + 46, cw - 22, card_h - 50, body, size=8, color=INK, spacing=1.2,
         name=f"Carte {lab} – texte")

# ---------- écosystème ----------
y += card_h + 12
text(s, M, y, CW, 40, [[
    ("EDLC s’appuie sur un écosystème engagé : ", {}), ("France Assureurs", B), (", ", {}),
    ("Bpifrance", B), (", ", {}), ("l’Agefiph", B), (", des réseaux d’accompagnement comme ", {}),
    ("l’Adie", B), (", ", {}), ("France Active", B), (", ", {}), ("BGE", B), (" ou ", {}),
    ("France Travail", B), (", ainsi que des assureurs partenaires. Cette coopération réunit les compétences "
                             "nécessaires autour d’un objectif commun : ", {}),
    ("pérenniser les micro-entreprises nouvellement créées.", B)]],
    size=8.6, spacing=1.25, name="Écosystème")

# ---------- impact social ----------
y += 48
text(s, M, y, CW, 20, "Un impact social documenté", size=14, font=TITLE_FONT, bold=True, color=BLUE,
     spacing=1.0, name="Titre – impact social")
y += 24
tile_w, tile_h, tg = 118, 53, 8
stats = [("58 %", "de femmes", BLUE, BLUE_T), ("82 %", "d’anciens demandeurs d’emploi", CORAL, CORAL_T),
         ("81 %", "de travailleurs handicapés", ORANGE, ORANGE_T), ("48 ans", "d’âge moyen", GREEN, GREEN_T)]
for i, (num, lab, c, t_) in enumerate(stats):
    r, col = divmod(i, 2)
    x = M + col * (tile_w + tg)
    yy = y + r * (tile_h + tg)
    rect(s, x, yy, tile_w, tile_h, t_, radius=10, name=f"Profil {num} – fond")
    text(s, x, yy + 8, tile_w, 24, num, size=18, font=TITLE_FONT, bold=True, color=c,
         align=PP_ALIGN.CENTER, spacing=1.0, name=f"Profil {num}")
    text(s, x + 4, yy + 32, tile_w - 8, 14, lab, size=7.4, color=INK, align=PP_ALIGN.CENTER,
         spacing=1.0, name=f"Profil {num} – légende")

# témoignage
qx = M + 2 * (tile_w + tg)
qw = CW - (qx - M)
qh = 2 * tile_h + tg
rect(s, qx, y, qw, qh, BLUE_T, radius=12, name="Témoignage – fond")
text(s, qx + 12, y + 2, 30, 34, "“", size=40, font=TITLE_FONT, bold=True, color=CORAL, spacing=1.0,
     name="Témoignage – guillemet (décor)")
oval(s, qx + qw - 68, y + 28, 56, YELLOW, name="Témoignage – cercle photo")
picture(s, f"{IMG}/amalia_raynal.png", qx + qw - 65.5, y + 30.5, w=51, name="Photo – Amalia Raynal")
text(s, qx + 14, y + 23, qw - 86, 70,
     "“En tant qu’entrepreneure, et d’autant plus dans un contexte de handicap ou de RQTH, la protection "
     "sociale peut être extrêment complexe. Grâce à EDLC, j’ai pu assurer mon activité et avancer "
     "sereinement dans mon coeur de métier, sans craindre l’imprévu. Avoir l’esprit tranquille est une clé "
     "essentielle et EDLC y répond !”",
     size=7.2, bold=True, italic=True, color=BLUE, spacing=1.15, name="Témoignage – citation")
text(s, qx + 14, y + qh - 15, qw - 24, 11, "Amalia Raynal, Assistante administrative à domicile",
     size=7.2, color=CORAL, bold=True, name="Témoignage – auteur")

# citation du président
y += qh + 8
lq_h = 52
rect(s, M, y, CW, lq_h, BLUE, radius=12, name="Citation président – fond")
text(s, M + 14, y + 4, 36, 40, "“", size=44, font=TITLE_FONT, bold=True, color=YELLOW, spacing=1.0,
     name="Citation président – guillemet (décor)")
text(s, M + 54, y + 9, CW - 70, 26,
     "“Dès son lancement Entrepreneurs de la Cité a dépassé les théories économiques conventionnelles pour "
     "ouvrir les portes de l’assurance aux créateurs d’entreprise les plus fragiles.”",
     size=8.2, bold=True, italic=True, color=WHITE, spacing=1.2, name="Citation président – texte")
text(s, M + 54, y + 35, CW - 70, 11, "Laurent DES BREST, Président de la Fondation EDLC", size=7.4,
     color=YELLOW, bold=True, name="Citation président – auteur")

# ---------- pied de page ----------
hline(s, M, H - 30, CW, CORAL, 0.75, name="Pied de page – filet")
text(s, M, H - 23, CW, 10, [[("ENTREPRENEURS DE LA CITÉ", {"bold": True, "color": CORAL}),
                             ("  •  Prévenir, protéger, pérenniser", {})]],
     size=6.8, color=GREY, align=PP_ALIGN.CENTER, spc=0.3, name="Pied de page")
text(s, W - M - 20, H - 23, 20, 10, "1", size=6.8, color=GREY, bold=True, align=PP_ALIGN.RIGHT,
     name="Numéro de page")

for sl in prs.slides:
    for el in list(sl.shapes._spTree.iter(qn("p:style"))):
        el.getparent().remove(el)
prs.save(OUT)
print("ok", OUT, "bas du contenu :", round(y + lq_h, 1), "/ pied", round(H - 30, 1))
