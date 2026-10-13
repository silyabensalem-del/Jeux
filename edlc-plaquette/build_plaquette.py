"""Génère la plaquette mécénat EDLC (2 pages A4) en PPTX avec éléments natifs, importable dans Canva."""
# Usage : python3 build_plaquette.py illustrations EDLC_Plaquette_mecenat.pptx
import sys
from PIL import Image
from pptx import Presentation
from pptx.util import Pt, Mm, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.oxml.ns import qn

IMG = sys.argv[1]
OUT = sys.argv[2]

# Charte EDLC
CORAL = RGBColor(0xEF, 0x3F, 0x34)
ORANGE = RGBColor(0xF8, 0x94, 0x10)
YELLOW = RGBColor(0xFF, 0xCB, 0x1F)
BLUE = RGBColor(0x26, 0x40, 0x8F)
GREEN = RGBColor(0x36, 0xB9, 0x95)
PINK = RGBColor(0xF5, 0x97, 0xB8)
INK = RGBColor(0x2B, 0x2B, 0x2B)
GREY = RGBColor(0x6B, 0x6B, 0x6B)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
# teintes claires dérivées des couleurs de la charte
CORAL_T = RGBColor(0xFD, 0xE8, 0xE6)
BLUE_T = RGBColor(0xE6, 0xEA, 0xF5)
YELLOW_T = RGBColor(0xFF, 0xF5, 0xD2)
ORANGE_T = RGBColor(0xFE, 0xEF, 0xDB)

TITLE_FONT = "Poppins"
BODY_FONT = "Open Sans"

W, H = 595.3, 841.9  # A4 en points
M = 36  # marge
CW = W - 2 * M

prs = Presentation()
prs.slide_width = Mm(210)
prs.slide_height = Mm(297)
blank = prs.slide_layouts[6]


def no_line(shp):
    shp.line.fill.background()


def rect(sl, x, y, w, h, fill, radius=None, line=None, lw=0.75, name=None):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    s = sl.shapes.add_shape(kind, Pt(x), Pt(y), Pt(w), Pt(h))
    if radius:
        s.adjustments[0] = min(0.5, radius / min(w, h))
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    if line is None:
        no_line(s)
    else:
        s.line.color.rgb = line
        s.line.width = Pt(lw)
    s.shadow.inherit = False
    s.text_frame.text = ""
    if name:
        s.name = name
    return s


def oval(sl, x, y, d, fill, name=None):
    s = sl.shapes.add_shape(MSO_SHAPE.OVAL, Pt(x), Pt(y), Pt(d), Pt(d))
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    no_line(s)
    s.shadow.inherit = False
    if name:
        s.name = name
    return s


def hline(sl, x, y, w, color, weight=0.75, name=None):
    c = sl.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Pt(x), Pt(y), Pt(x + w), Pt(y))
    c.line.color.rgb = color
    c.line.width = Pt(weight)
    if name:
        c.name = name
    return c


def text(sl, x, y, w, h, paras, size=9, color=INK, font=BODY_FONT, bold=False, italic=False,
         align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, spacing=1.2, after=0, name=None, spc=None):
    """paras: str | list de paragraphes ; un paragraphe = str | list de runs (texte, dict d'options)."""
    tb = sl.shapes.add_textbox(Pt(x), Pt(y), Pt(w), Pt(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    if isinstance(paras, str):
        paras = [paras]
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = spacing
        p.space_after = Pt(after)
        runs = [(para, {})] if isinstance(para, str) else [(r, {}) if isinstance(r, str) else r for r in para]
        for t, o in runs:
            r = p.add_run()
            r.text = t
            f = r.font
            f.name = o.get("font", font)
            f.size = Pt(o.get("size", size))
            f.bold = o.get("bold", bold)
            f.italic = o.get("italic", italic)
            f.color.rgb = o.get("color", color)
            s_ = o.get("spc", spc)
            if s_ is not None:
                r._r.get_or_add_rPr().set("spc", str(int(s_ * 100)))
            if "link" in o:
                r.hyperlink.address = o["link"]
                f.color.rgb = o.get("color", color)
                f.underline = True
    if name:
        tb.name = name
    return tb


def picture(sl, path, x, y, w=None, h=None, name=None):
    im = Image.open(path)
    bbox = im.getchannel("A").getbbox()
    if bbox and bbox != (0, 0) + im.size:
        im = im.crop(bbox)
        path = path.replace(".png", "_crop.png")
        im.save(path)
    iw, ih = im.size
    if w and not h:
        h = w * ih / iw
    if h and not w:
        w = h * iw / ih
    p = sl.shapes.add_picture(path, Pt(x), Pt(y), Pt(w), Pt(h))
    if name:
        p.name = name
    return p, w, h


def group(sl, shapes, name):
    g = sl.shapes.add_group_shape(shapes)
    g.name = name
    return g


def section_title(sl, x, y, label, size=15, color=CORAL, bar=CORAL, w=CW, name=None):
    b = rect(sl, x, y + 3, 4.5, size + 2, bar, radius=2.2, name=(name or label) + " – repère")
    t = text(sl, x + 11, y, w - 11, size * 1.5, label, size=size, color=color, font=TITLE_FONT,
             bold=True, spacing=1.0, name=(name or label))
    return b, t


def sub_title(sl, x, y, label, w=CW, color=BLUE, size=11.5):
    return text(sl, x, y, w, size * 1.5, label, size=size, color=color, font=TITLE_FONT, bold=True,
                spacing=1.0, name=label)


def chrome(sl, page):
    """Bandeau haut, en-tête et pied de page (identiques sur les 2 pages)."""
    rect(sl, 0, 0, W, 7, CORAL, name="Bandeau haut corail")
    text(sl, W - M - 300, 15, 300, 10, "ENTREPRENEURS DE LA CITÉ  |  FONDATION RECONNUE D’UTILITÉ PUBLIQUE",
         size=6.5, color=GREY, align=PP_ALIGN.RIGHT, spc=0.4, name="En-tête")
    hline(sl, M, H - 34, CW, CORAL, 0.75, name="Pied de page – filet")
    text(sl, M, H - 27, CW, 10, [[("ENTREPRENEURS DE LA CITÉ", {"bold": True, "color": CORAL}),
                                  ("  •  Prévenir, protéger, pérenniser", {})]],
         size=6.8, color=GREY, align=PP_ALIGN.CENTER, spc=0.3, name="Pied de page")
    text(sl, W - M - 20, H - 27, 20, 10, str(page), size=6.8, color=GREY, bold=True,
         align=PP_ALIGN.RIGHT, name="Numéro de page")


# ---------- pictogrammes vectoriels (formes natives) ----------
def preset(sl, kind, x, y, w, h, fill=None, line=None, lw=0.9, name=None, rot=0):
    sh = sl.shapes.add_shape(kind, Pt(x), Pt(y), Pt(w), Pt(h))
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid(); sh.fill.fore_color.rgb = fill
    if line is None:
        no_line(sh)
    else:
        sh.line.color.rgb = line; sh.line.width = Pt(lw)
    sh.rotation = rot
    if name:
        sh.name = name
    return sh


def icon_protect(sl, x, y, d, bg):
    """Bouclier + coche (formes natives, non groupées pour rester vectorielles dans Canva)."""
    oval(sl, x, y, d, bg, name="Picto protéger – fond")
    s = d / 30
    preset(sl, MSO_SHAPE.FLOWCHART_OFFPAGE_CONNECTOR, x + 8 * s, y + 6.5 * s, 14 * s, 17 * s, WHITE,
           name="Picto protéger – bouclier")
    # coche = deux barres arrondies pivotées
    preset(sl, MSO_SHAPE.ROUNDED_RECTANGLE, x + 10.2 * s, y + 13.4 * s, 5 * s, 2.2 * s, bg, rot=45,
           name="Picto protéger – coche (court)")
    preset(sl, MSO_SHAPE.ROUNDED_RECTANGLE, x + 12.6 * s, y + 11.8 * s, 8.5 * s, 2.2 * s, bg, rot=-50,
           name="Picto protéger – coche (long)")


def icon_prevent(sl, x, y, d, bg, fg):
    """Ampoule = prévention / formation."""
    oval(sl, x, y, d, bg, name="Picto prévenir – fond")
    s = d / 30
    oval(sl, x + 9 * s, y + 5 * s, 12 * s, fg, name="Picto prévenir – ampoule")
    rect(sl, x + 12 * s, y + 15.5 * s, 6 * s, 5 * s, fg, name="Picto prévenir – culot")
    rect(sl, x + 12 * s, y + 21.3 * s, 6 * s, 1.4 * s, fg, name="Picto prévenir – trait 1")
    rect(sl, x + 13 * s, y + 23.5 * s, 4 * s, 1.4 * s, fg, name="Picto prévenir – trait 2")


def icon_mail(sl, x, y, w, color):
    h = w * 0.68
    rect(sl, x, y, w, h, None, line=color, lw=0.9, name="Picto e-mail – enveloppe")
    preset(sl, MSO_SHAPE.FLOWCHART_MERGE, x, y, w, h * 0.55, None, line=color, name="Picto e-mail – rabat")


def icon_pin(sl, x, y, d, color, hole):
    preset(sl, MSO_SHAPE.FLOWCHART_MERGE, x + d * 0.14, y + d * 0.55, d * 0.72, d * 0.75, color,
           name="Picto adresse – pointe")
    oval(sl, x, y, d, color, name="Picto adresse – tête")
    oval(sl, x + d * 0.3, y + d * 0.3, d * 0.4, hole, name="Picto adresse – centre")


def icon_web(sl, x, y, d, color):
    preset(sl, MSO_SHAPE.OVAL, x, y, d, d, None, line=color, name="Picto web – globe")
    preset(sl, MSO_SHAPE.OVAL, x + d * 0.3, y, d * 0.4, d, None, line=color, name="Picto web – méridien")
    hline(sl, x, y + d / 2, d, color, 0.9, name="Picto web – équateur")


# =====================================================================
# PAGE 1
# =====================================================================
s1 = prs.slides.add_slide(blank)
chrome(s1, 1)

picture(s1, f"{IMG}/logo_edlc.png", M, 26, w=128, name="Logo EDLC")

# --- Héro ---
blob = oval(s1, 418, 62, 150, YELLOW_T, name="Décor – cercle jaune")
picture(s1, f"{IMG}/entrepreneur_maison.png", 425, 72, w=142, name="Illustration – entrepreneur et maison")

text(s1, M, 96, 360, 12, "ENTREPRENEURS DE LA CITÉ", size=9, bold=True, color=CORAL, spc=1.5,
     name="Surtitre")
text(s1, M, 111, 375, 92, "L’assurance sociale et solidaire au service de l’entrepreneuriat",
     size=23, font=TITLE_FONT, bold=True, color=CORAL, spacing=1.0, name="Titre principal")
text(s1, M, 180, 370, 32,
     "Devenez mécène d’une économie où chacun peut entreprendre — et faire durer son activité.",
     size=11, color=BLUE, bold=True, spacing=1.15, name="Sous-titre")

# piliers
px = M
for label, bg, fg in (("PRÉVENIR", YELLOW, BLUE), ("PROTÉGER", CORAL, WHITE), ("PÉRENNISER", BLUE, WHITE)):
    wpill = 9 + len(label) * 5.6
    pill = rect(s1, px, 220, wpill, 17, bg, radius=8.5, name=f"Pilier {label} – fond")
    t = text(s1, px, 220, wpill, 17, label, size=7.5, bold=True, color=fg, align=PP_ALIGN.CENTER,
             anchor=MSO_ANCHOR.MIDDLE, spc=0.8, name=f"Pilier {label}")
    group(s1, [pill, t], f"Pilier {label}")
    px += wpill + 4
    if label != "PÉRENNISER":
        text(s1, px, 220, 8, 17, "•", size=9, bold=True, color=CORAL, align=PP_ALIGN.CENTER,
             anchor=MSO_ANCHOR.MIDDLE, name="Séparateur piliers")
        px += 12

# --- Introduction ---
rect(s1, M, 252, 3, 40, CORAL, name="Introduction – filet")
text(s1, M + 11, 251, CW - 11, 44,
     "Depuis 2007, la Fondation Entrepreneurs de la Cité (EDLC) ouvre l’accès à une assurance adaptée aux "
     "créateurs d’entreprise en situation de vulnérabilité. Nous agissons là où les offres classiques "
     "répondent mal aux besoins : au croisement de l’inclusion, de l’emploi et de la pérennité économique.",
     size=9, spacing=1.25, name="Introduction")

# --- Chiffres clés ---
y = 306
tw = (CW - 3 * 8) / 4
row1 = [("3 362", "entrepreneurs assurés\nfin 2025", CORAL, CORAL_T),
        ("2 544", "personnes formées aux risques\nen 2025", BLUE, BLUE_T),
        ("89 %", "de survie après sinistre\n(données 2025)", ORANGE, ORANGE_T),
        ("350+", "activités\ncouvertes", GREEN, RGBColor(0xE1, 0xF5, 0xEF))]
for i, (num, lab, c, t_) in enumerate(row1):
    x = M + i * (tw + 8)
    bg = rect(s1, x, y, tw, 54, t_, radius=7, name=f"Chiffre {num} – fond")
    n = text(s1, x, y + 6, tw, 24, num, size=19, font=TITLE_FONT, bold=True, color=c,
             align=PP_ALIGN.CENTER, spacing=1.0, name=f"Chiffre {num}")
    l = text(s1, x + 4, y + 31, tw - 8, 20, lab.split("\n"), size=7.2, color=INK, align=PP_ALIGN.CENTER,
             spacing=1.05, name=f"Chiffre {num} – légende")
    group(s1, [bg, n, l], f"Chiffre clé {num}")
y2 = y + 62
tw3 = (CW - 2 * 8) / 3
row2 = [("11 741", "Entrepreneurs assures depuis 2007"),
        ("38 816", "personnes formées aux risques depuis 2007"),
        ("20 ans", "d’expérience aux côtés des entrepreneurs")]
for i, (num, lab) in enumerate(row2):
    x = M + i * (tw3 + 8)
    bg = rect(s1, x, y2, tw3, 40, BLUE, radius=7, name=f"Chiffre {num} – fond")
    n = text(s1, x + 12, y2, 66, 40, num, size=17, font=TITLE_FONT, bold=True, color=YELLOW,
             anchor=MSO_ANCHOR.MIDDLE, spacing=1.0, name=f"Chiffre {num}")
    l = text(s1, x + 80, y2, tw3 - 88, 40, lab, size=7.4, color=WHITE, anchor=MSO_ANCHOR.MIDDLE,
             spacing=1.1, name=f"Chiffre {num} – légende")
    group(s1, [bg, n, l], f"Chiffre clé {num}")

# --- Pourquoi soutenir EDLC ? ---
y = 428
section_title(s1, M, y, "Pourquoi soutenir EDLC ?", size=15)
text(s1, M, y + 27, CW, 80, [
    "Les créateurs en situation de précarité ou de handicap peuvent rencontrer des difficultés à accéder aux "
    "assurances professionnelles classiques. Sans protection adaptée, un accident, un litige ou un arrêt de "
    "travail peut compromettre l’activité qu’ils ont créée.",
    "EDLC apporte une réponse concrète et complémentaire à l’accompagnement à la création d’entreprise : pour "
    "bénéficier gratuitement de nos services, les créateurs sont accompagnés par une structure d’aide à la "
    "création."], size=9, spacing=1.25, after=5, name="Pourquoi soutenir – texte")

# --- Un modèle qui associe protection et prévention ---
y = 538
sub_title(s1, M, y, "Un modèle qui associe protection et prévention")
cw2 = (CW - 10) / 2
cards = [("PROTÉGER", CORAL, CORAL_T,
          [[("Des solutions assurantielles à tarifs solidaires, couvrant notamment la multirisque professionnelle, "
             "la responsabilité civile professionnelle, la santé, la prévoyance et la responsabilité civile "
             "décennale. ", {"bold": True}),
            ("L’assistance juridique est proposée gratuitement.", {})]]),
         ("PRÉVENIR", BLUE, YELLOW_T,
          [[("Des webinaires hebdomadaires, des formations et des interventions auprès des réseaux "
             "d’accompagnement pour aider les entrepreneurs à identifier les risques et à choisir les "
             "protections adaptées.", {"bold": True})]])]
for i, (lab, c, t_, body) in enumerate(cards):
    x = M + i * (cw2 + 10)
    yc = y + 21
    rect(s1, x, yc, cw2, 104, t_, radius=9, name=f"Carte {lab} – fond")
    if lab == "PROTÉGER":
        icon_protect(s1, x + 10, yc + 10, 28, CORAL)
    else:
        icon_prevent(s1, x + 10, yc + 10, 28, BLUE, YELLOW)
    text(s1, x + 46, yc + 10, cw2 - 56, 14, lab, size=10, font=TITLE_FONT, bold=True, color=c, spc=0.6,
         spacing=1.0, name=f"Carte {lab} – titre")
    text(s1, x + 46, yc + 27, cw2 - 56, 74, body, size=8.1, color=INK, spacing=1.18,
         name=f"Carte {lab} – texte")

# --- Un impact social documenté ---
y = 676
sub_title(s1, M, y, "Un impact social documenté")
stats = [("58", " % de femmes", CORAL), ("82", " % d’anciens demandeurs d’emploi", ORANGE),
         ("81", " % de travailleurs handicapés", BLUE), (None, "Moyenne d’âge de 48 ans", GREEN)]
sw = (CW - 3 * 8) / 4
for i, (num, lab, c) in enumerate(stats):
    x = M + i * (sw + 8)
    yc = y + 21
    bg = rect(s1, x, yc, sw, 46, WHITE, radius=7, line=c, lw=1.1, name=f"Profil {i+1} – cadre")
    acc = rect(s1, x, yc + 9, 3.5, 28, c, name=f"Profil {i+1} – accent")
    if num:
        runs = [[(num, {"font": TITLE_FONT, "size": 15, "color": c})], [(lab.strip(), {"size": 7.6})]]
    else:
        runs = [[("Moyenne d’âge de", {"size": 7.6})], [("48 ans", {"font": TITLE_FONT, "size": 15, "color": c})]]
    t = text(s1, x + 10, yc, sw - 14, 46, runs, bold=True, color=INK, anchor=MSO_ANCHOR.MIDDLE,
             spacing=1.0, name=f"Profil {i+1}")
    group(s1, [bg, acc, t], f"Profil bénéficiaires {i+1}")

text(s1, M, y + 78, CW, 50,
     "EDLC s’appuie sur un écosystème engagé : France Assureurs, Bpifrance (Cap Créa), l’Agefiph, des réseaux "
     "d’accompagnement comme l’Adie, France Active, BGE ou France Travail, ainsi que des assureurs partenaires. "
     "Cette coopération réunit les compétences nécessaires autour d’un objectif commun : pérenniser les "
     "micro-entreprises nouvellement créées.", size=8.6, spacing=1.22, name="Écosystème – texte")

# =====================================================================
# PAGE 2
# =====================================================================
s2 = prs.slides.add_slide(blank)
chrome(s2, 2)
picture(s2, f"{IMG}/logo_edlc.png", M, 26, w=110, name="Logo EDLC")

y = 82
section_title(s2, M, y, "Votre don, un levier d’action concret", size=19)
text(s2, M, y + 31, CW, 40,
     "En devenant mécène, votre entreprise contribue à financer la solidarité au cœur de notre modèle : pas de "
     "frais de gestion ni de commission sur la gestion des solutions, des protections accessibles et des "
     "actions de prévention au plus près des entrepreneurs.", size=9, spacing=1.25, name="Levier – texte")

# --- Impact concret : 4 cartes ---
y = 161
sub_title(s2, M, y, "Quel est l’impact concret de votre soutien ?")
impacts = [
    ("5 000 €", YELLOW, BLUE, YELLOW_T,
     [[("30 entrepreneurs", {"bold": True}), (" accèdent à l’assurance sociale et solidaire et sont accompagnés pendant 1 an", {})]]),
    ("10 000 €", ORANGE, WHITE, ORANGE_T,
     [[("60 entrepreneurs", {"bold": True}), (" accèdent à l’assurance sociale et solidaire ", {})], "ou",
      [("44 entrepreneurs", {"bold": True}), (" sont formés à la gestion des risques liés à l’entrepreneuriat", {})]]),
    ("20 000 €", CORAL, WHITE, CORAL_T,
     [[("120 entrepreneurs", {"bold": True}), (" accèdent à l’assurance sociale et solidaire ", {})], "ou",
      [("88 entrepreneurs", {"bold": True}), (" sont formés à la gestion des risques", {})]]),
    ("50 000 €", BLUE, WHITE, BLUE_T,
     [[("300 entrepreneurs", {"bold": True}), (" accèdent à l’assurance sociale et solidaire ", {})], "ou",
      [("220 entrepreneurs", {"bold": True}), (" sont formés à la gestion des risques", {})]]),
]
gap = 9
cw4 = (CW - 3 * gap) / 4
yc = y + 22
ch = 164
for i, (amt, c, fg, t_, body) in enumerate(impacts):
    x = M + i * (cw4 + gap)
    rect(s2, x, yc, cw4, ch, t_, radius=9, line=c, lw=1, name=f"Carte don {amt} – fond")
    rect(s2, x, yc, cw4, 46, c, radius=9, name=f"Carte don {amt} – bandeau")
    rect(s2, x, yc + 30, cw4, 16, c, name=f"Carte don {amt} – bandeau (raccord)")
    text(s2, x + 10, yc + 7, cw4 - 20, 10, "DON", size=6.8, bold=True, color=fg, spc=1.2, name=f"Carte don {amt} – libellé")
    text(s2, x + 10, yc + 17, cw4 - 20, 24, amt, size=17, font=TITLE_FONT, bold=True, color=fg, spacing=1.0,
         name=f"Carte don {amt} – montant")
    text(s2, x + 10, yc + 54, cw4 - 20, 10, "IMPACT MESURABLE", size=6.3, bold=True,
         color=(ORANGE if c == YELLOW else c), spc=1.0, name=f"Carte don {amt} – libellé impact")
    paras = []
    for p in body:
        if p == "ou":
            paras.append([("ou", {"bold": True, "italic": True, "color": (ORANGE if c == YELLOW else c)})])
        else:
            paras.append(p)
    text(s2, x + 10, yc + 67, cw4 - 16, ch - 70, paras, size=7.8, color=INK, spacing=1.15, after=3,
         name=f"Carte don {amt} – impact")

# --- Coût net ---
y = yc + ch + 14
sub_title(s2, M, y, "Quel est le coût net de votre mécénat ?")
text(s2, M, y + 20, CW, 26,
     "À titre indicatif, pour une entreprise éligible au régime français du mécénat, la réduction d’impôt est de "
     "60 % du don dans la limite des plafonds légaux applicables*.", size=9, spacing=1.25, name="Coût net – texte")

ty = y + 54
tbl_w = 372
rows = [("DON", "RÉDUCTION D’IMPÔT\nINDICATIVE (60 %)", "COÛT NET\nINDICATIF"),
        ("5 000 €", "3 000 €", "2 000 €"), ("10 000 €", "6 000 €", "4 000 €"),
        ("20 000 €", "12 000 €", "8 000 €"), ("50 000 €", "30 000 €", "20 000 €")]
gt = s2.shapes.add_table(5, 3, Pt(M), Pt(ty), Pt(tbl_w), Pt(30 + 4 * 19.5))
gt.name = "Tableau coût net du mécénat"
tbl = gt.table
# supprime le style de tableau par défaut (bandes)
tblPr = tbl._tbl.tblPr
tblPr.set("firstRow", "0"); tblPr.set("bandRow", "0")
for j, wcol in enumerate((110, 142, 120)):
    tbl.columns[j].width = Pt(wcol)
tbl.rows[0].height = Pt(30)
for r in range(1, 5):
    tbl.rows[r].height = Pt(19.5)
for r, row in enumerate(rows):
    for j, val in enumerate(row):
        cell = tbl.cell(r, j)
        cell.margin_left = cell.margin_right = Pt(6)
        cell.margin_top = cell.margin_bottom = Pt(2)
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        cell.fill.solid()
        if r == 0:
            cell.fill.fore_color.rgb = BLUE
        elif j == 2:
            cell.fill.fore_color.rgb = CORAL_T
        else:
            cell.fill.fore_color.rgb = WHITE if r % 2 else RGBColor(0xF6, 0xF7, 0xFB)
        tf = cell.text_frame
        tf.word_wrap = True
        lines = val.split("\n")
        for k, ln in enumerate(lines):
            p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
            p.alignment = PP_ALIGN.CENTER
            p.line_spacing = 1.0
            run = p.add_run(); run.text = ln
            f = run.font; f.name = BODY_FONT
            if r == 0:
                f.size = Pt(7.2); f.bold = True; f.color.rgb = WHITE
            else:
                f.size = Pt(9 if j != 1 else 8.8)
                f.bold = j != 1
                f.color.rgb = CORAL if j == 2 else (BLUE if j == 0 else INK)
        # bordures fines
        tcPr = cell._tc.get_or_add_tcPr()
        for edge in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
            ln = tcPr.makeelement(qn(edge), {"w": str(Pt(0.5)), "cap": "flat", "cmpd": "sng", "algn": "ctr"})
            sf = ln.makeelement(qn("a:solidFill"), {})
            clr = sf.makeelement(qn("a:srgbClr"), {"val": "D9DCE6"})
            sf.append(clr); ln.append(sf)
            tcPr.insert(0, ln) if False else tcPr.append(ln)

picture(s2, f"{IMG}/analyse_document.png", M + tbl_w + 14, ty - 2, w=CW - tbl_w - 14, name="Illustration – analyse du document")

# note fiscale
ny = ty + 30 + 4 * 19.5 + 9
text(s2, M, ny, CW, 50,
     "*À noter : le taux et le montant réellement imputable dépendent de la situation fiscale du mécène et des "
     "plafonds applicables. En règle générale, les dons ouvrant droit à réduction sont retenus dans la limite de "
     "20 000 € ou de 0,5 % du chiffre d’affaires hors taxes si ce montant est plus élevé. Pour les versements "
     "dépassant 2 M€ sur l’année, le taux de réduction est généralement de 40 % sur la fraction excédentaire. "
     "La réduction peut être reportée dans les conditions prévues par la loi. Faites confirmer votre situation "
     "par votre conseil fiscal.", size=6.9, italic=True, color=GREY, spacing=1.2, name="Note fiscale")

# engagement (pastille)
ey = ny + 48
rect(s2, M, ey, CW, 24, YELLOW_T, radius=12, name="Engagement – fond")
oval(s2, M + 12, ey + 9, 6, CORAL, name="Engagement – puce")
text(s2, M + 26, ey, CW - 36, 24,
     "Votre engagement peut prendre la forme d’un don annuel ou pluriannuel au service de l’inclusion "
     "entrepreneuriale.", size=8.6, bold=True, color=BLUE, anchor=MSO_ANCHOR.MIDDLE, name="Engagement – texte")

# --- Contact ---
cy = ey + 32
chh = H - 44 - cy
rect(s2, M, cy, CW, chh, BLUE, radius=12, name="Contact – fond bleu")
oval(s2, W - M - 128, cy + chh - 120, 112, RGBColor(0x33, 0x50, 0xA3), name="Décor – cercle contact")
picture(s2, f"{IMG}/interlocutrice.png", W - M - 132, cy + chh - 112, w=118, name="Illustration – interlocutrice")
cx = M + 18
text(s2, cx, cy + 14, 360, 24, "Engageons-nous ensemble", size=17, font=TITLE_FONT, bold=True, color=WHITE,
     spacing=1.0, name="Contact – titre")
text(s2, cx, cy + 42, 400, 26,
     [[("Pour toute information, n’hésitez pas à contacter ", {}),
       ("Claire BENISTI - Déléguée Générale", {"bold": True}), (" :", {})]],
     size=8.8, color=WHITE, spacing=1.2, name="Contact – texte")
icon_mail(s2, cx, cy + 75, 11, YELLOW)
text(s2, cx + 17, cy + 72, 330, 13,
     [[("claire.benisti@entrepreneursdelacite.org",
        {"link": "mailto:claire.benisti@entrepreneursdelacite.org", "color": YELLOW, "bold": True})]],
     size=8.8, name="Contact – e-mail")
hline(s2, cx, cy + 93, 340, RGBColor(0x5A, 0x70, 0xB5), 0.6, name="Contact – séparateur")
icon_pin(s2, cx + 1, cy + 104, 8.5, YELLOW, BLUE)
text(s2, cx + 17, cy + 101, 330, 28,
     [[("Fondation Entrepreneurs de la Cité", {"bold": True})],
      "31 place Jules Grandclément — 69100 Villeurbanne"],
     size=8.4, color=WHITE, spacing=1.2, name="Contact – adresse")
icon_web(s2, cx, cy + 133, 10, YELLOW)
text(s2, cx + 17, cy + 132, 330, 13,
     [[("www.entrepreneursdelacite.org", {"link": "https://www.entrepreneursdelacite.org", "color": WHITE})]],
     size=8.4, color=WHITE, name="Contact – site web")
text(s2, cx, cy + 153, 350, 28,
     "Prenons contact pour construire un partenariat à la hauteur de vos engagements.",
     size=9.2, bold=True, color=YELLOW, spacing=1.15, name="Contact – accroche")

for sl in prs.slides:
    for el in sl.shapes._spTree.iter(qn("p:style")):
        pass
    for el in list(sl.shapes._spTree.iter(qn("p:style"))):
        el.getparent().remove(el)
prs.save(OUT)
print("ok", OUT, "contact block height", chh)
